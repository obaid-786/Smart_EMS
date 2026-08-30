# -*- coding: utf-8 -*-
"""Generate Word explanation + prompt for Smart EMS & MIS backup."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Smart_EMS_MIS_Step_by_Step_Explanation.docx")

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")

def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")

# Title
title = doc.add_heading("Smart EMS & MIS — Step-by-Step Explanation", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("SUGNA SPONGE AND POWER PVT LTD\nBackup package · Coding by Amazad Ali · Version 1.0")
r.italic = True

p("This Word file explains how the system works, which files are required, how to run it, "
  "and includes a ready AI prompt you can use to recreate or extend this code.")

# 1
h("1. What this project is")
p("Smart EMS & MIS is a plant monitoring system for Continuous Casting Machine (CCM) and Rolling Mill. "
  "It reads live data from a Siemens S7-1200 PLC (PAC3200 energy meter + billet/HMD counts), "
  "shows Live / EMS / MIS screens, saves heats and energy to a chosen folder, and exports Excel reports.")
bullets([
    "Classic HMI: http://localhost:5000/  (plc_live_monitor.html + plant_mis.js)",
    "Pro UI: http://localhost:5000/pro/  (steel_ems_pro/)",
    "Bridge API: bridge_server.py (Flask) polls PLC every 1 second and saves data",
])

# 2
h("2. Files included in this backup (required)")
h("2.1 Root / bridge", 2)
bullets([
    "bridge_server.py — main server: PLC poll, CSV/JSONL log, APIs (/api/live, /api/energy, /api/bands, production heats, folder config)",
    "pac3200_reader.py — Snap7 PLC tag map (DB4 energy + DB8 billets/HMD)",
    "requirements.txt — Python packages (flask, flask-cors, python-snap7, openpyxl)",
    "log_config.json — selected Plant_Data folder path",
    "plc_live_monitor.html — Classic HMI",
    "plant_mis.js — Classic MIS / heat logic",
    "mes_dashboard.js — MES dashboard helpers",
    "mes/ — SQLite MES package (auth, energy, production, reports)",
])
h("2.2 Pro UI (steel_ems_pro/)", 2)
bullets([
    "index.html — navigation (Plant, CCM, Rolling Mill, EMS, MIS, Smart Detector, Health)",
    "assets/pro.js — casting timer, shared heat, EMS 24h energy, Smart Detector, folder chooser, Excel export",
    "assets/pro.css — layout and report table styles",
])
h("2.3 Data saved at runtime (not in this zip of code — created when bridge runs)", 2)
bullets([
    "pac3200_log.csv — 1-second energy + HMD log",
    "plant_live.jsonl — full live snapshots",
    "ccm_heats.jsonl — CCM heat end records (casting time, strands, tons)",
    "rm_heats.jsonl — Rolling Mill heat records",
    "ems_readings.jsonl — EMS period / month energy readings",
    "plant_mes.sqlite3 — MES database",
])

# 3
h("3. How to install and run (step by step)")
numbered([
    "Install Python 3.10+ on the PC that will talk to the PLC.",
    "Copy this whole backup folder to a working path, e.g. D:\\Smart EMS and MIS Project",
    "Open Command Prompt in that folder.",
    "Create venv (optional): python -m venv .venv  then  .venv\\Scripts\\activate",
    "Install packages: pip install -r requirements.txt",
    "Edit pac3200_reader.py if PLC IP / DB offsets differ for your plant.",
    "Start bridge FROM PROJECT ROOT (not from steel_ems_pro):  python bridge_server.py",
    "Open browser: Classic http://localhost:5000/   Pro http://localhost:5000/pro/",
    "In Pro → CCM Live or Health: choose Save folder (e.g. D:\\Smart EMS and MIS Project\\Plant_Data) → Save folder.",
    "Keep bridge running — data auto-saves every 1 second to that folder.",
])
p("Important: Always run bridge_server.py from the project root so /pro/ and APIs load correctly.", bold=True)

# 4
h("4. Feature explanation — CCM casting time")
numbered([
    "When casting is running, KPI shows Casting time (running clock) and auto-end countdown (~75 minutes).",
    "When heat stops (manual End or auto 75 min), casting time + heat data are saved to ccm_heats.jsonl in the selected folder BEFORE the timer resets.",
    "While stopped, KPI shows Next heat tentative (plant cycle ~2.3 hours from heat start), not the casting clock.",
    "When next heat starts, casting running time resets to 00:00:00.",
    "Heat number is shared with Rolling Mill via browser localStorage key steel_ems_shared_heat.",
])

# 5
h("5. Feature explanation — Rolling Mill EMS energy")
numbered([
    "Window presets: Plant day (e.g. 09:00 → next day 09:00), previous plant day, last 24h, today, or custom From/To.",
    "Calculate & save loads /api/energy (meter ΔkWh / ΔkVAh / ΔkVArh + calculated ∫kW/kVA/kVAr) and /api/bands (on-load/idle/stop).",
    "Each calculation auto-saves into ems_readings.jsonl in the selected folder.",
    "Month Excel downloads day-wise rows for the whole month with meter start/end/delta and calculated energy; each day is also auto-saved.",
])

# 6
h("6. Feature explanation — data folder chooser")
numbered([
    "Pro UI (CCM Live and Health) and Classic HMI both use POST /api/log-config with { log_dir: \"D\\\\...\\\\Plant_Data\" }.",
    "Bridge tests write access, saves path to log_config.json, and redirects all CSV/JSONL/SQLite writes there.",
    "Export all (ZIP) uses /api/data/export.zip.",
    "Browser cannot open a Windows folder dialog for the server disk — you type the path on the PC running the bridge.",
])

# 7
h("7. Feature explanation — MIS report tables")
p("Rolling Mill heats table is full width with larger sticky headers (Heat, Ended, R3, TMT, Miss, t, On-load, Idle, Util%, kWh/t) for clearer focus.")

# 8
h("8. Main API list (bridge_server.py)")
bullets([
    "GET /api/live — latest PLC snapshot",
    "GET /api/history?start&end — CSV rows",
    "GET /api/energy?start&end — meter + calculated energy for a period",
    "GET /api/energy/days?year&month&day_start_hour — day-wise 24h plant days",
    "GET /api/energy/month.xls?year&month&day_start_hour — month Excel download",
    "GET /api/bands?start&end&start_kw&onload_kw — on-load / idle / stop productivity",
    "GET/POST /api/log-config — choose / read save folder",
    "GET/POST /api/production/ccm_heats|rm_heats|ems_readings — save/load heats & EMS readings",
    "GET /api/data/export.zip — backup all plant data files",
])

# 9
h("9. Department navigation (Pro)")
bullets([
    "Plant home — company brand + overview",
    "CCM — Live (cast), EMS, MIS",
    "Rolling Mill — Live (mill timers/HMD), EMS kWh, MIS, Smart Detector",
    "CPP / 100 DRI — Coming soon placeholders",
    "Health & trust — PLC status + data folder chooser",
])

# 10
h("10. AI PROMPT — copy this to rebuild or extend the code")
p("Copy the block below into Cursor / ChatGPT when you want another agent to recreate or continue this product:", bold=True)

prompt = r'''
You are building / extending "Smart EMS & MIS" for SUGNA SPONGE AND POWER PVT LTD.

STACK
- Python Flask bridge_server.py on port 5000, run from project ROOT only.
- pac3200_reader.py (python-snap7) reads S7-1200 DB4 PAC3200 + DB8 billets/HMD every 1s.
- Classic HMI at / from plc_live_monitor.html + plant_mis.js (keep safe; do not break).
- Pro UI at /pro/ from steel_ems_pro/ (index.html, assets/pro.js, assets/pro.css) — additive only.
- Persist all plant data into a user-selected folder via GET/POST /api/log-config → log_config.json.
- Files in that folder: pac3200_log.csv, plant_live.jsonl, ccm_heats.jsonl, rm_heats.jsonl, ems_readings.jsonl, plant_mes.sqlite3.

CCM CASTING RULES
- While casting: show Casting time (running) + auto end ~75 min.
- On heat end (manual or auto): SAVE castSec + heat data to ccm_heats.jsonl BEFORE reset.
- When stopped: show Next heat tentative (cycle ~2.3 h from startedAt), not casting clock.
- On next heat start: reset casting timer; share heatNo with Rolling Mill via localStorage key steel_ems_shared_heat.

ROLLING MILL EMS
- Presets for 24h windows (plant day start hour default 9 → next day 9).
- Calculate meter ΔkWh/ΔkVAh/ΔkVArh + calculated ∫kW/kVA/kVAr; auto-save to ems_readings.jsonl.
- Month Excel: one row per plant day with window start/end, meter start/end/delta, calculated kWh/kVAh/kVArh.
- APIs: /api/energy, /api/energy/days, /api/energy/month.xls, /api/bands.

UI RULES
- Sidebar brand: only "Smart EMS & MIS".
- Departments accordion (CCM, Rolling Mill, CPP, 100 DRI) start minimized.
- CCM Live has NO mill on-load/idle/stop (those belong to Rolling Mill).
- MIS RM heats table full width with strong header focus.
- Pro data folder chooser on CCM Live + Health; top bar shows Save path.
- Bump ?v= cache on pro.js/pro.css/index.html after edits.

DELIVER
- Keep classic HMI working.
- Sync changes to delivery copy if present.
- Explain how to run: pip install -r requirements.txt && python bridge_server.py
'''

box = doc.add_paragraph()
run = box.add_run(prompt.strip())
run.font.name = "Consolas"
run.font.size = Pt(9)

# 11
h("11. Checklist after restore")
numbered([
    "pip install -r requirements.txt",
    "python bridge_server.py from project root",
    "Open /pro/ and hard refresh (Ctrl+F5)",
    "Set Save folder and click Save folder",
    "Confirm files appear in Plant_Data (CSV growing every second)",
    "Test CCM Start/End heat → ccm_heats.jsonl grows",
    "Test Rolling Mill EMS Calculate & save + Month Excel",
])

h("12. Support notes")
bullets([
    "If /pro/ 404 or old UI: stop old bridge and restart from the folder that contains steel_ems_pro.",
    "If PLC offline: UI still opens; live tags show offline until Snap7 connects.",
    "Default MES users (if enabled): see MES_README.md (admin/admin123 etc.).",
    "Coding by Amazad Ali · Version 1.0",
])

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
er = end.add_run("— End of explanation —")
er.italic = True

doc.save(OUT)
print("Wrote", OUT)
